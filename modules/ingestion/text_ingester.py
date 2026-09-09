import os
from modules.embeddings.text_embedding import chunk_text


def ingest_text(file_path: str, upload_id: str, original_filename: str = None) -> tuple[list[dict], dict]:
    """
    Process plain text, Markdown, or Word documents (.txt, .md, .docx).
    """
    if original_filename is None:
        original_filename = os.path.basename(file_path)

    ext = os.path.splitext(file_path)[1].lower()
    full_text = ""

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        lines.append(" | ".join(row_cells))
            full_text = "\n".join(lines)
        except Exception as e:
            print(f"[TEXT INGESTER] Error reading DOCX: {e}")
    else:
        # .txt or .md
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read().strip()

    if not full_text:
        full_text = f"Empty text document {original_filename}."

    raw_chunks = chunk_text(full_text)
    if not raw_chunks:
        raw_chunks = [full_text]

    chunks = []
    for i, chk in enumerate(raw_chunks):
        chunks.append({
            "text": chk,
            "upload_id": upload_id,
            "source_file": original_filename,
            "content_type": "text",
            "chunk_id": i + 1
        })

    doc_meta = {
        "upload_id": upload_id,
        "source_file": original_filename,
        "content_type": "text",
        "chunks": len(chunks),
        "status": "indexed"
    }

    return chunks, doc_meta
