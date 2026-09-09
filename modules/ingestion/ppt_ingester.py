import os
from io import BytesIO
from PIL import Image
from pptx import Presentation
from modules.ocr.ocr_engine import run_ocr
from modules.embeddings.text_embedding import chunk_text
from backend.config import MAX_PPT_SLIDES


def ingest_pptx(file_path: str, upload_id: str, original_filename: str = None) -> tuple[list[dict], dict]:
    """
    Process a PowerPoint presentation (.pptx) slide by slide:
    1. Extract slide shapes and table text.
    2. Extract speaker notes.
    3. Extract embedded pictures and run Tesseract OCR on them.
    4. Create slide-aware chunks preserving slide numbers.
    """
    if original_filename is None:
        original_filename = os.path.basename(file_path)

    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    if total_slides == 0:
        raise ValueError(f"Presentation '{original_filename}' has 0 slides.")

    slides_to_process = min(total_slides, MAX_PPT_SLIDES)
    chunks = []
    chunk_counter = 0

    for slide_idx in range(slides_to_process):
        slide_num = slide_idx + 1
        slide = prs.slides[slide_idx]
        slide_texts = []

        # 1. Text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        slide_texts.append(" | ".join(row_cells))
            # Embedded pictures OCR
            elif shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    image_bytes = shape.image.blob
                    pil_img = Image.open(BytesIO(image_bytes))
                    # Only OCR reasonably sized images
                    if pil_img.width > 100 and pil_img.height > 100:
                        ocr_res = run_ocr(pil_img)
                        ocr_words = [str(item[1]) for item in ocr_res if len(item) >= 2]
                        if ocr_words:
                            slide_texts.append(f"[Image Text: {' '.join(ocr_words)}]")
                except Exception:
                    pass

        # 2. Speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_texts.append(f"[Speaker Notes: {notes}]")

        combined_slide_text = "\n".join(slide_texts).strip()
        if not combined_slide_text:
            combined_slide_text = f"[Slide {slide_num} of {original_filename} with visual or graphic content]"

        slide_chunks = chunk_text(combined_slide_text)
        if not slide_chunks:
            slide_chunks = [combined_slide_text]

        for s_chk in slide_chunks:
            chunk_counter += 1
            chunks.append({
                "text": s_chk,
                "upload_id": upload_id,
                "source_file": original_filename,
                "content_type": "pptx",
                "slide": slide_num,
                "chunk_id": chunk_counter
            })

    doc_meta = {
        "upload_id": upload_id,
        "source_file": original_filename,
        "content_type": "pptx",
        "slides": slides_to_process,
        "total_slides": total_slides,
        "chunks": len(chunks),
        "status": "indexed"
    }

    return chunks, doc_meta
